"""Export helper for Data Absensi Finger.

Produces Excel and PDF from the same ABSENSI-final row dictionaries so
the two downloads contain the same dataset.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, LongTable, TableStyle, Paragraph
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


HEADERS = [
    ("No.", "no"), ("Hari", "hari"), ("Tgl Kerja", "tgl_kerja"),
    ("FingerID", "finger_id"), ("Nama", "nama"),
    ("Baku (IN)", "jam_baku_in"), ("Baku (OUT)", "jam_baku_out"),
    ("IN", "jam_in"), ("OUT", "jam_out"),
    ("Awal TLM", "awal_tlm"), ("Total TLM", "total_tlm"),
    ("TK TLM", "tingkat_tlm"), ("Pot TLM (%)", "persen_pot_tlm"),
    ("Awal PSW", "awal_psw"), ("Total PSW", "total_psw"),
    ("TK PSW", "tingkat_psw"), ("Pot PSW (%)", "persen_pot_psw"),
    ("KET", "ket"),
]

HEADER_FILL = "172033"
GRID = "D5DBE3"
ZEBRA = "EEF2F7"


def _value(row, key):
    value = row.get(key)
    return "" if value is None else str(value)


def build_excel(rows, period_label=""):
    output = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Data Absensi"

    ws["A1"] = "DATA ABSENSI FINGER - HASIL EKSPOR"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))
    ws["A2"] = "Periode"
    ws["B2"] = period_label
    ws["A2"].font = Font(bold=True)

    header_row = 4
    for col, (label, _) in enumerate(HEADERS, 1):
        cell = ws.cell(header_row, col, label)
        cell.fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, row in enumerate(rows, header_row + 1):
        for col, (_, key) in enumerate(HEADERS, 1):
            cell = ws.cell(row_index, col, _value(row, key))
            cell.alignment = Alignment(vertical="center")

        ket_color = str(row.get("ket_color") or "").strip()
        if ket_color:
            for col in range(1, len(HEADERS) + 1):
                cell = ws.cell(row_index, col)
                cell.fill = PatternFill(
                    start_color="FF" + ket_color,
                    end_color="FF" + ket_color,
                    fill_type="solid",
                )
                cell.font = Font(
                    color="FFFFFFFF",
                    bold=(col == len(HEADERS)),
                )
        elif row_index % 2 == 0:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row_index, col).fill = PatternFill(
                    start_color=ZEBRA,
                    end_color=ZEBRA,
                    fill_type="solid",
                )

    widths = [7, 12, 16, 13, 30, 12, 12, 12, 12, 12, 12, 12, 14, 12, 12, 12, 14, 12]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(HEADERS))}{max(header_row, header_row + len(rows))}"
    ws.sheet_view.showGridLines = False

    wb.save(output)
    output.seek(0)
    return output


def build_pdf(rows, period_label=""):
    output = BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("DataAbsensiTitle", parent=styles["Heading2"],
                                 fontName="Helvetica-Bold", fontSize=11, leading=14, spaceAfter=3)
    meta_style = ParagraphStyle("DataAbsensiMeta", parent=styles["Normal"],
                                fontSize=7, leading=9, spaceAfter=5)
    cell_style = ParagraphStyle("DataAbsensiCell", parent=styles["Normal"],
                                fontSize=5.8, leading=7)
    head_style = ParagraphStyle("DataAbsensiHead", parent=cell_style,
                                fontName="Helvetica-Bold", textColor=colors.white, alignment=1)

    document = SimpleDocTemplate(
        output, pagesize=landscape(A4),
        rightMargin=7 * mm, leftMargin=7 * mm, topMargin=7 * mm, bottomMargin=7 * mm,
    )

    table_data = [[Paragraph(label, head_style) for label, _ in HEADERS]]
    for row in rows:
        table_data.append([Paragraph(_value(row, key), cell_style) for _, key in HEADERS])

    col_widths = [
        8 * mm, 14 * mm, 18 * mm, 15 * mm, 32 * mm,
        14 * mm, 14 * mm, 14 * mm, 14 * mm,
        14 * mm, 14 * mm, 14 * mm, 16 * mm,
        14 * mm, 14 * mm, 14 * mm, 16 * mm, 14 * mm,
    ]

    table = LongTable(table_data, colWidths=col_widths, repeatRows=1)
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_FILL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#" + GRID)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (0, -1), "CENTER"),
        ("ALIGN", (3, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]
    for row_index, row in enumerate(rows, 1):
        ket_color = str(row.get("ket_color") or "").strip()
        if ket_color:
            commands.extend([
                (
                    "BACKGROUND",
                    (0, row_index),
                    (-1, row_index),
                    colors.HexColor("#" + ket_color),
                ),
                (
                    "TEXTCOLOR",
                    (0, row_index),
                    (-1, row_index),
                    colors.white,
                ),
            ])
        elif row_index % 2 == 0:
            commands.append((
                "BACKGROUND",
                (0, row_index),
                (-1, row_index),
                colors.HexColor("#" + ZEBRA),
            ))
    table.setStyle(TableStyle(commands))

    story = [
        Paragraph("DATA ABSENSI FINGER - HASIL EKSPOR", title_style),
        Paragraph(f"Periode: {period_label}", meta_style),
        table,
    ]
    document.build(story)
    output.seek(0)
    return output

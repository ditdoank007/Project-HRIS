"""
Global report export presentation standard for HRIS Reborn.

Single source of truth for report colors and legend semantics.
PDF and Excel exporters should consume these definitions instead of
declaring their own color palettes.
"""

from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors


REPORT_COLORS = {
    # Header ini adalah warna sumber tunggal untuk PDF, Excel,
    # dan laporan lain yang memakai standar export global.
    "header": "172033",
    "header_text": "FFFFFF",
    "holiday": "B91C1C",
    "siaga": "15803D",
    "blue": "2563EB",
    "orange": "B45309",
    "wfh": "475569",
    "normal": "172033",
    "zebra": "EEF2F7",
    "grid": "D5DBE3",
}


def ket_color_key(code, status_um=None, master_potongan_codes=None):
    """
    Global color controller untuk KET pada seluruh report HRIS Reborn.

    Semantik:
      - DL/OP/SD + StatusUM=1  -> orange (potong Uang Makan)
      - DL/OP/SD + StatusUM!=1 -> blue (tidak potong Uang Makan)
      - shift1/shift2          -> siaga (hijau)
      - WFH                    -> wfh (abu-abu gelap)
      - kode Master Potongan ketidakhadiran yang memotong UM -> orange

    Return value adalah key REPORT_COLORS agar UI/Excel/PDF memakai
    satu palette global.
    """
    normalized = str(code or '').strip().upper()
    if not normalized:
        return 'normal'

    if normalized in ('SHIFT1', 'SHIFT2'):
        return 'siaga'

    if normalized == 'WFH':
        return 'wfh'

    if normalized in ('DL', 'OP', 'SD'):
        try:
            return 'orange' if int(status_um or 0) == 1 else 'blue'
        except (TypeError, ValueError):
            return 'blue'

    # Kode ketidakhadiran yang memang memotong Uang Makan.
    # Termasuk seluruh kode yang berasal dari Master Potongan.
    fixed_cut_codes = {
        'CT', 'CAP', 'S', 'S-1', 'S-2', 'I', 'IJIN', 'IZIN', 'ALPA'
    }
    master_codes = {
        str(value or '').strip().upper()
        for value in (master_potongan_codes or set())
        if str(value or '').strip()
    }

    if normalized in fixed_cut_codes or normalized in master_codes:
        return 'orange'

    return 'normal'


REPORT_LEGEND = (
    ("holiday", "HARI LIBUR"),
    ("siaga", "SIAGA"),
    ("blue", "SPRIN/DL TIDAK MEMOTONG UANG MAKAN"),
    ("orange", "SPRIN/DL MEMOTONG UANG MAKAN"),
    ("wfh", "ABSEN ONLINE WFH"),
)


# Standard nama hari untuk seluruh format laporan HRIS Reborn.
HARI_INDONESIA = {
    0: "SEN",
    1: "SEL",
    2: "RAB",
    3: "KAM",
    4: "JUM",
    5: "SAB",
    6: "MIN",
}


def format_hari_indonesia(tanggal):
    """Return nama hari laporan dalam Bahasa Indonesia."""
    if not tanggal:
        return ""
    return HARI_INDONESIA.get(tanggal.weekday(), "")


def report_header_day_color(tanggal):
    """Return warna teks nama hari pada header PDF laporan."""
    if not tanggal:
        return REPORT_COLORS["header_text"]
    if tanggal.weekday() >= 5:
        return REPORT_COLORS["holiday"]
    return REPORT_COLORS["header_text"]


def excel_font_color(color_key):
    """Return an OpenPyXL-compatible ARGB font color."""
    return "FF" + REPORT_COLORS.get(color_key, REPORT_COLORS["normal"])


def excel_header_fill():
    """Return the standard report header fill for Excel exports."""
    return excel_fill("header")


def excel_header_font():
    """Return the standard white header font for Excel exports."""
    return Font(
        color="FFFFFFFF",
        bold=True,
    )


def excel_fill(color_key):
    """Return a solid OpenPyXL fill using the global report palette."""
    value = REPORT_COLORS.get(color_key, REPORT_COLORS["normal"])
    return PatternFill(
        start_color=value,
        end_color=value,
        fill_type="solid",
    )


def pdf_color(color_key):
    """Return a ReportLab color using the global report palette."""
    value = REPORT_COLORS.get(color_key, REPORT_COLORS["normal"])
    return colors.HexColor("#" + value)


def add_excel_legend(ws, start_row, label_column=3, color_column=2):
    """Render the standard report legend in an Excel worksheet."""
    row = start_row

    for color_key, label in REPORT_LEGEND:
        color_cell = ws.cell(row=row, column=color_column)
        color_cell.fill = excel_fill(color_key)
        color_cell.font = Font(
            color="FFFFFFFF",
            bold=True,
        )

        label_cell = ws.cell(row=row, column=label_column, value=label)
        label_cell.font = Font(bold=True)
        row += 1

    return row


def pdf_legend_text():
    """Return the standard legend as colored HTML for ReportLab."""
    parts = []

    for color_key, label in REPORT_LEGEND:
        parts.append(
            '<font color="#{color}"><b>{label}</b></font>'.format(
                color=REPORT_COLORS[color_key],
                label=label,
            )
        )

    return "Keterangan: " + " &nbsp;&nbsp; ".join(parts) + "."
